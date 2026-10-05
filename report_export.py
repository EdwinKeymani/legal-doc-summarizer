"""
Export a document's analysis as a Word (.docx) or PDF report.

Covers proposal requirement 4.2: "export summaries as downloadable PDF or
Word reports".

Both formats are built from the same ReportContent, so they always contain
the same information:
  title and details -> summary -> flagged clauses -> key provisions -> key terms

The AI summary arrives as sanitised HTML (the same Markdown -> nh3 path the
web page uses). It is converted into simple blocks (heading / paragraph /
bullet, each with bold and italic runs) that both writers understand.
"""

import io
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from html import unescape
from html.parser import HTMLParser

DISCLAIMER = (
    "This report was generated automatically. Summaries may be produced by an AI model "
    "and clauses are found by keyword rules, so items can be missed or misclassified. "
    "It is informational only and is not legal advice."
)

SEVERITY_COLOURS = {  # (text colour, background colour) as hex, matching the app
    "High": ("B91C1C", "FEF2F2"),
    "Medium": ("B45309", "FFFBEB"),
    "Low": ("15803D", "F0FDF4"),
    "Review": ("64748B", "F1F5F9"),
}


# ---------------------------------------------------------------------------
# Shared report structure
# ---------------------------------------------------------------------------
@dataclass
class ReportContent:
    title: str
    details: list                      # [(label, value)]
    summary_blocks: list               # [(block_type, [(text, bold, italic)])]
    clauses: list                      # [{"category", "text", "severity", "explanation"}]
    provision_sections: list           # [(plural label, description, [(party, [texts])])]
    key_terms: dict                    # {term_type: [{"value", "context", "date"}]}
    notes: list = field(default_factory=list)
    generated_at: str = ""
    brand_name: str = "Chambua"


class _SummaryHtmlParser(HTMLParser):
    """Turns the limited, sanitised summary HTML into blocks of styled runs."""

    BLOCK_TAGS = {"p": "paragraph", "h1": "heading", "h2": "heading", "h3": "heading", "h4": "heading", "blockquote": "paragraph"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self.current_runs = None
        self.current_type = None
        self.bold_depth = 0
        self.italic_depth = 0
        self.list_stack = []  # "ul" or "ol"
        self.ordered_counters = []

    def _start_block(self, block_type):
        self._end_block()
        self.current_type, self.current_runs = block_type, []

    def _end_block(self):
        if self.current_runs is not None:
            runs = [(text, bold, italic) for text, bold, italic in self.current_runs if text]
            if runs and "".join(text for text, _, _ in runs).strip():
                first_text = runs[0][0].lstrip()
                runs[0] = (first_text, runs[0][1], runs[0][2])
                self.blocks.append((self.current_type, runs))
        self.current_runs, self.current_type = None, None

    def handle_starttag(self, tag, attrs):
        if tag in self.BLOCK_TAGS and not self.list_stack:
            self._start_block(self.BLOCK_TAGS[tag])
        elif tag in ("ul", "ol"):
            self._end_block()
            self.list_stack.append(tag)
            self.ordered_counters.append(0)
        elif tag == "li":
            if self.list_stack and self.list_stack[-1] == "ol":
                self.ordered_counters[-1] += 1
                self._start_block("numbered")
                self.current_runs.append((f"{self.ordered_counters[-1]}. ", False, False))
            else:
                self._start_block("bullet")
        elif tag in ("strong", "b"):
            self.bold_depth += 1
        elif tag in ("em", "i"):
            self.italic_depth += 1
        elif tag == "br" and self.current_runs is not None:
            self.current_runs.append(("\n", False, False))

    def handle_endtag(self, tag):
        if tag in self.BLOCK_TAGS and not self.list_stack:
            self._end_block()
        elif tag == "li":
            self._end_block()
        elif tag in ("ul", "ol"):
            self._end_block()
            if self.list_stack:
                self.list_stack.pop()
                self.ordered_counters.pop()
        elif tag in ("strong", "b"):
            self.bold_depth = max(0, self.bold_depth - 1)
        elif tag in ("em", "i"):
            self.italic_depth = max(0, self.italic_depth - 1)

    def handle_data(self, data):
        if self.current_runs is None:
            if not data.strip():
                return
            self._start_block("paragraph")
        text = re.sub(r"\s+", " ", data)
        self.current_runs.append((text, self.bold_depth > 0, self.italic_depth > 0))

    def close(self):
        super().close()
        self._end_block()


def summary_html_to_blocks(safe_html):
    parser = _SummaryHtmlParser()
    parser.feed(str(safe_html or ""))
    parser.close()
    return parser.blocks


BRAND_LOGO_PNG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "brand", "logo-mark.png")


def build_report_content(document, safe_summary_html, provision_groups, key_terms, risk_level_note=None,
                         brand_name="Chambua"):
    """Collect everything both formats need from one document.
    provision_groups: [(type, description, [(party, [texts])])]"""
    analysis = document.analysis
    upload_time = document.upload_date
    if upload_time is not None and upload_time.tzinfo is None:
        upload_time = upload_time.replace(tzinfo=timezone.utc)

    details = [
        ("Document", document.file_name),
        ("Format", document.file_format or "-"),
        ("Uploaded", upload_time.strftime("%d %b %Y, %H:%M UTC") if upload_time else "-"),
        ("Status", document.status or "-"),
    ]
    if analysis:
        details += [
            ("Summary method", analysis.summary_method or "-"),
            ("Options chosen", " / ".join(filter(None, [analysis.summary_depth, analysis.extraction_mode,
                                                       analysis.risk_level, analysis.context_hint]))),
        ]
        if analysis.processing_seconds is not None:
            details.append(("Processing time", f"{analysis.processing_seconds:.1f} seconds"))

    notes = []
    if analysis and analysis.truncated_for_ai:
        notes.append("The AI summary covers only the first ~30,000 characters of this long document.")
    if risk_level_note:
        notes.append(risk_level_note)

    return ReportContent(
        title=document.file_name,
        details=details,
        summary_blocks=summary_html_to_blocks(safe_summary_html) or [("paragraph", [("No summary available.", False, True)])],
        clauses=[{"category": clause.clause_category or "-", "text": clause.extracted_text or "",
                  "severity": clause.risk_severity or "Review", "explanation": clause.explanation or ""}
                 for clause in document.clauses],
        provision_sections=provision_groups,
        key_terms=key_terms,
        notes=notes,
        generated_at=datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC"),
        brand_name=brand_name,
    )


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------
def export_docx(content):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    word_document = Document()
    for section in word_document.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2)
        # Header: logo mark + wordmark, on every page
        header_paragraph = section.header.paragraphs[0]
        if os.path.exists(BRAND_LOGO_PNG):
            header_paragraph.add_run().add_picture(BRAND_LOGO_PNG, width=Cm(0.6))
        name_run = header_paragraph.add_run("  " + content.brand_name.lower())
        name_run.bold = True
        name_run.font.size = Pt(12)
        name_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        dot_run = header_paragraph.add_run(".")
        dot_run.bold = True
        dot_run.font.size = Pt(12)
        dot_run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.text = f"{content.brand_name} report, generated {content.generated_at}. Informational only, not legal advice."
        footer_paragraph.runs[0].font.size = Pt(8)
        footer_paragraph.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    normal_style = word_document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    def shade_cell(cell, hex_colour):
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), hex_colour)
        properties.append(shading)

    def add_runs(paragraph, runs):
        for text, bold, italic in runs:
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic

    def make_table(headers, rows, column_widths_cm):
        table = word_document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = ""
            run = cell.paragraphs[0].add_run(header)
            run.bold = True
            run.font.size = Pt(9)
            shade_cell(cell, "F1F5F9")
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = ""
                run = cells[index].paragraphs[0].add_run(str(value))
                run.font.size = Pt(9)
        # Fixed layout plus the table's own column grid: Word and LibreOffice
        # otherwise ignore per-cell widths and make every column equal.
        table.autofit = False
        for index, width in enumerate(column_widths_cm):
            table.columns[index].width = Cm(width)
        for row in table.rows:
            for index, width in enumerate(column_widths_cm):
                row.cells[index].width = Cm(width)
        return table

    title = word_document.add_heading("Document Analysis Report", level=0)
    title.runs[0].font.size = Pt(22)
    subtitle = word_document.add_paragraph()
    subtitle_run = subtitle.add_run(content.title)
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(13)

    make_table(["Detail", "Value"], content.details, [4.0, 12.6])
    for note in content.notes:
        note_paragraph = word_document.add_paragraph()
        note_run = note_paragraph.add_run("Note: " + note)
        note_run.italic = True
        note_run.font.size = Pt(9)
    disclaimer_paragraph = word_document.add_paragraph()
    disclaimer_run = disclaimer_paragraph.add_run(DISCLAIMER)
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(8.5)
    disclaimer_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    word_document.add_heading("Summary", level=1)
    for block_type, runs in content.summary_blocks:
        if block_type == "heading":
            add_runs(word_document.add_heading(level=2), runs)
        elif block_type == "bullet":
            add_runs(word_document.add_paragraph(style="List Bullet"), runs)
        else:  # paragraph, numbered (number already in the text)
            add_runs(word_document.add_paragraph(), runs)

    word_document.add_heading(f"Flagged Clauses ({len(content.clauses)})", level=1)
    if content.clauses:
        table = make_table(["Category", "Extracted text", "Risk", "Assessment"],
                           [(clause["category"], clause["text"], clause["severity"], clause["explanation"])
                            for clause in content.clauses], [3.0, 6.6, 1.8, 5.2])
        for row_index, clause in enumerate(content.clauses, start=1):
            text_colour, background_colour = SEVERITY_COLOURS.get(clause["severity"], SEVERITY_COLOURS["Review"])
            risk_cell = table.rows[row_index].cells[2]
            shade_cell(risk_cell, background_colour)
            risk_run = risk_cell.paragraphs[0].runs[0]
            risk_run.bold = True
            risk_run.font.color.rgb = RGBColor.from_string(text_colour)
    else:
        word_document.add_paragraph("No clauses were flagged.")

    word_document.add_heading("Key Provisions", level=1)
    intro = word_document.add_paragraph()
    intro_run = intro.add_run("Found by keyword rules. Grouped by the party the sentence names as its subject.")
    intro_run.italic = True
    intro_run.font.size = Pt(9)
    if not any(groups for _, _, groups in content.provision_sections):
        word_document.add_paragraph("No provisions found, or the document text is not stored (uploaded before 3 October 2026).")
    for provision_label, description, groups in content.provision_sections:
        if not groups:
            continue
        total = sum(len(texts) for _, texts in groups)
        word_document.add_heading(f"{provision_label} ({total}): {description}", level=2)
        for party, texts in groups:
            party_paragraph = word_document.add_paragraph()
            party_run = party_paragraph.add_run(party)
            party_run.bold = True
            for text in texts:
                word_document.add_paragraph(text, style="List Bullet")

    word_document.add_heading("Key Dates, Deadlines, Amounts and Rates", level=1)
    term_rows = []
    for term_type, items in content.key_terms.items():
        for item in items:
            shown_value = item["value"] + (f"  ({item['date']:%d %b %Y})" if item.get("date") else "")
            term_rows.append((term_type, shown_value, item["context"]))
    if term_rows:
        make_table(["Type", "Value", "In context"], term_rows, [3.0, 4.2, 9.4])
    else:
        word_document.add_paragraph("No dates, deadlines, amounts or rates found.")

    output = io.BytesIO()
    word_document.save(output)
    return output.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_safe(text):
    """ReportLab's built-in fonts cover the Windows-1252 character set, which
    includes curly quotes, dashes, bullets and the section sign. Anything else
    is replaced with '?' rather than printing as an empty box. Then escape the
    characters ReportLab's paragraph markup treats specially."""
    text = str(text or "").encode("cp1252", errors="replace").decode("cp1252")
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def export_pdf(content):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    base_styles = getSampleStyleSheet()
    ink, muted, rule = colors.HexColor("#0F172A"), colors.HexColor("#64748B"), colors.HexColor("#E2E8F0")
    styles = {
        "title": ParagraphStyle("ReportTitle", parent=base_styles["Title"], fontSize=20, leading=24, alignment=TA_LEFT, textColor=ink, spaceAfter=2),
        "subtitle": ParagraphStyle("ReportSubtitle", parent=base_styles["Normal"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=ink, spaceAfter=10),
        "h1": ParagraphStyle("H1", parent=base_styles["Heading2"], fontSize=14, leading=18, textColor=colors.HexColor("#15803D"), spaceBefore=14, spaceAfter=6),
        "h2": ParagraphStyle("H2", parent=base_styles["Heading3"], fontSize=11.5, leading=15, textColor=ink, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("Body", parent=base_styles["Normal"], fontSize=10, leading=14.5, textColor=ink, spaceAfter=5),
        "small": ParagraphStyle("Small", parent=base_styles["Normal"], fontSize=8.5, leading=11.5, textColor=ink),
        "small_bold": ParagraphStyle("SmallBold", parent=base_styles["Normal"], fontName="Helvetica-Bold", fontSize=8.5, leading=11.5, textColor=ink),
        "note": ParagraphStyle("Note", parent=base_styles["Italic"], fontSize=8.5, leading=12, textColor=muted, spaceAfter=4),
        "party": ParagraphStyle("Party", parent=base_styles["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4, spaceAfter=2, textColor=ink),
    }

    def runs_to_markup(runs):
        markup = ""
        for text, bold, italic in runs:
            piece = _pdf_safe(text).replace("\n", "<br/>")
            if bold:
                piece = f"<b>{piece}</b>"
            if italic:
                piece = f"<i>{piece}</i>"
            markup += piece
        return markup

    def cell(text, style="small"):
        return Paragraph(_pdf_safe(text), styles[style])

    def table(rows, column_widths, header=True, extra_styles=()):
        result = Table(rows, colWidths=column_widths, repeatRows=1 if header else 0)
        commands = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.5, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        if header:
            commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")))
        result.setStyle(TableStyle(commands + list(extra_styles)))
        return result

    usable_width = A4[0] - 4.4 * cm
    story = [
        Paragraph("Document Analysis Report", styles["title"]),
        Paragraph(_pdf_safe(content.title), styles["subtitle"]),
        table([[cell(label, "small_bold"), cell(value)] for label, value in content.details],
              [3.6 * cm, usable_width - 3.6 * cm], header=False),
        Spacer(1, 6),
    ]
    story += [Paragraph("Note: " + _pdf_safe(note), styles["note"]) for note in content.notes]
    story.append(Paragraph(_pdf_safe(DISCLAIMER), styles["note"]))

    story.append(Paragraph("Summary", styles["h1"]))
    pending_bullets = []

    def flush_bullets():
        if pending_bullets:
            story.append(ListFlowable([ListItem(Paragraph(markup, styles["body"]), leftIndent=12) for markup in pending_bullets],
                                      bulletType="bullet", start="\u2022", leftIndent=14, bulletFontSize=8))
            pending_bullets.clear()

    for block_type, runs in content.summary_blocks:
        if block_type == "bullet":
            pending_bullets.append(runs_to_markup(runs))
            continue
        flush_bullets()
        story.append(Paragraph(runs_to_markup(runs), styles["h2" if block_type == "heading" else "body"]))
    flush_bullets()

    story.append(Paragraph(f"Flagged Clauses ({len(content.clauses)})", styles["h1"]))
    if content.clauses:
        rows = [[cell("Category", "small_bold"), cell("Extracted text", "small_bold"), cell("Risk", "small_bold"), cell("Assessment", "small_bold")]]
        severity_styles = []
        for row_index, clause in enumerate(content.clauses, start=1):
            text_colour, background_colour = SEVERITY_COLOURS.get(clause["severity"], SEVERITY_COLOURS["Review"])
            risk_style = ParagraphStyle(f"Risk{row_index}", parent=styles["small_bold"], textColor=colors.HexColor("#" + text_colour))
            rows.append([cell(clause["category"]), cell(clause["text"]), Paragraph(_pdf_safe(clause["severity"]), risk_style), cell(clause["explanation"])])
            severity_styles.append(("BACKGROUND", (2, row_index), (2, row_index), colors.HexColor("#" + background_colour)))
        story.append(table(rows, [2.8 * cm, 6.0 * cm, 1.7 * cm, usable_width - 10.5 * cm], extra_styles=severity_styles))
    else:
        story.append(Paragraph("No clauses were flagged.", styles["body"]))

    story.append(Paragraph("Key Provisions", styles["h1"]))
    story.append(Paragraph("Found by keyword rules. Grouped by the party the sentence names as its subject.", styles["note"]))
    if not any(groups for _, _, groups in content.provision_sections):
        story.append(Paragraph("No provisions found, or the document text is not stored (uploaded before 3 October 2026).", styles["body"]))
    for provision_label, description, groups in content.provision_sections:
        if not groups:
            continue
        total = sum(len(texts) for _, texts in groups)
        story.append(Paragraph(f"{provision_label} ({total}): {_pdf_safe(description)}", styles["h2"]))
        for party, texts in groups:
            items = [ListItem(Paragraph(_pdf_safe(text), styles["small"]), leftIndent=12) for text in texts]
            story.append(KeepTogether([Paragraph(_pdf_safe(party), styles["party"]),
                                       ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=14, bulletFontSize=7)])
                         if len(texts) <= 4 else Paragraph(_pdf_safe(party), styles["party"]))
            if len(texts) > 4:
                story.append(ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=14, bulletFontSize=7))

    story.append(Paragraph("Key Dates, Deadlines, Amounts and Rates", styles["h1"]))
    rows = [[cell("Type", "small_bold"), cell("Value", "small_bold"), cell("In context", "small_bold")]]
    for term_type, items in content.key_terms.items():
        for item in items:
            shown_value = item["value"] + (f" ({item['date']:%d %b %Y})" if item.get("date") else "")
            rows.append([cell(term_type), cell(shown_value), cell(item["context"])])
    if len(rows) > 1:
        story.append(table(rows, [2.8 * cm, 3.8 * cm, usable_width - 6.6 * cm]))
    else:
        story.append(Paragraph("No dates, deadlines, amounts or rates found.", styles["body"]))

    def draw_brand_mark(canvas, left, top, size):
        """The Chambua page-and-bars mark as vector shapes (same geometry as
        the SVG logo, 64-unit grid, y flipped for PDF coordinates)."""
        unit = size / 64.0
        def point(x, y):
            return left + x * unit, top - y * unit
        canvas.saveState()
        canvas.setLineWidth(3.5 * unit)
        canvas.setLineJoin(1)
        canvas.setStrokeColor(colors.HexColor("#0F172A"))
        canvas.setFillColor(colors.white)
        page_outline = canvas.beginPath()
        page_outline.moveTo(*point(14, 6)); page_outline.lineTo(*point(40, 6)); page_outline.lineTo(*point(52, 18))
        page_outline.lineTo(*point(52, 56)); page_outline.lineTo(*point(48, 60)); page_outline.lineTo(*point(14, 60))
        page_outline.lineTo(*point(10, 56)); page_outline.lineTo(*point(10, 10)); page_outline.close()
        canvas.drawPath(page_outline, stroke=1, fill=1)
        fold = canvas.beginPath()
        fold.moveTo(*point(40, 6)); fold.lineTo(*point(40, 18)); fold.lineTo(*point(52, 18))
        canvas.drawPath(fold, stroke=1, fill=0)
        for x, y, width, colour in [(17, 25, 28, "#DC2626"), (17, 36, 22, "#D97706"), (17, 47, 26, "#16A34A")]:
            canvas.setFillColor(colors.HexColor(colour))
            bar_left, bar_top = point(x, y)
            canvas.roundRect(bar_left, bar_top - 6 * unit, width * unit, 6 * unit, 3 * unit, stroke=0, fill=1)
        canvas.restoreState()

    def draw_page_frame(canvas, pdf_document):
        canvas.saveState()
        # Header: logo mark + wordmark
        header_top = A4[1] - 0.9 * cm
        draw_brand_mark(canvas, 2.2 * cm, header_top, 0.62 * cm)
        canvas.setFont("Helvetica-Bold", 11)
        canvas.setFillColor(ink)
        wordmark = content.brand_name.lower()
        text_left = 2.2 * cm + 0.8 * cm
        canvas.drawString(text_left, header_top - 0.45 * cm, wordmark)
        canvas.setFillColor(colors.HexColor("#DC2626"))
        canvas.drawString(text_left + canvas.stringWidth(wordmark, "Helvetica-Bold", 11), header_top - 0.45 * cm, ".")
        # Footer
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(2.2 * cm, 1.2 * cm, f"{content.brand_name} report, generated {content.generated_at}. Informational only, not legal advice.")
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, f"Page {pdf_document.page}")
        canvas.restoreState()

    output = io.BytesIO()
    pdf_document = SimpleDocTemplate(output, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                                     topMargin=2.2 * cm, bottomMargin=2 * cm,
                                     title=f"Analysis report: {content.title}", author=content.brand_name)
    pdf_document.build(story, onFirstPage=draw_page_frame, onLaterPages=draw_page_frame)
    return output.getvalue()


def report_filename(original_file_name, extension):
    """'Service Agreement.docx' -> 'Service-Agreement-analysis-report.pdf'"""
    stem = original_file_name.rsplit(".", 1)[0]
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", unescape(stem)).strip("-") or "document"
    return f"{safe_stem[:80]}-analysis-report.{extension}"
